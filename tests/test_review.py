from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from journalprose.catalog import load
from journalprose.review import audit_text, mask_markdown, read_manuscript


def row(text, check, **kwargs):
    return next(c for c in audit_text(text, **kwargs)["checks"] if c["id"] == check)


class ReviewTests(unittest.TestCase):
    def test_empty_input(self):
        result = audit_text("")
        self.assertEqual(result["total_matches"], 0)
        self.assertEqual(result["check_count"], 21)
        self.assertNotIn("quality_score", result)

    def test_limit_preserves_total(self):
        result = row("not only " * 17, "contrast-not-only", max_per_check=2)
        self.assertEqual(result["total_matches"], 17)
        self.assertEqual(len(result["examples"]), 2)
        self.assertTrue(result["truncated"])

    def test_bad_limits(self):
        for value in (0, -1, 1001, True, "2"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                audit_text("x", max_per_check=value)

    def test_invalid_text(self):
        with self.assertRaises(ValueError):
            audit_text(None)

    def test_invalid_location_unit(self):
        with self.assertRaises(ValueError):
            audit_text("x", unit="page")

    def test_line_column_and_offsets(self):
        text = "Clean.\n  not only a record."
        finding = row(text, "contrast-not-only")["examples"][0]
        self.assertEqual((finding["line"], finding["column"]), (2, 3))
        self.assertEqual(text[finding["start"]:finding["end"]], "not only")

    def test_windows_newlines_normalized(self):
        finding = row("Clean.\r\nnot only data", "contrast-not-only")["examples"][0]
        self.assertEqual(finding["line"], 2)
        self.assertEqual(finding["column"], 1)

    def test_cross_line_contrast(self):
        self.assertEqual(row("The record is not\ncomplete but usable.", "contrast-not-but")["total_matches"], 1)

    def test_contrast_does_not_cross_sentence(self):
        self.assertEqual(row("The record is not complete. But it is usable.", "contrast-not-but")["total_matches"], 0)

    def test_necessary_negation_is_not_a_blanket_violation(self):
        result = audit_text("The records were not complete. Not all samples were retained. The estimate does not include late entries.")
        self.assertEqual(result["policy_matches"], 0)

    def test_clear_demonstrative_avoids_candidate(self):
        self.assertEqual(row("This restriction limits the sample.", "vague-reference")["total_matches"], 0)

    def test_code_fence_and_inline_code_excluded(self):
        text = "```text\nnot only\n```\n\nUse `rather than` as a search string.\n\nnot merely a record."
        self.assertEqual(row(text, "contrast-not-only")["total_matches"], 0)
        self.assertEqual(row(text, "contrast-rather")["total_matches"], 0)
        self.assertEqual(row(text, "contrast-not-merely")["total_matches"], 1)

    def test_tilde_fence_excluded(self):
        self.assertEqual(row("~~~\nnot only\n~~~\n", "contrast-not-only")["total_matches"], 0)

    def test_longer_fence_requires_matching_end(self):
        text = "````\n```\nnot only\n````\nnot only"
        self.assertEqual(row(text, "contrast-not-only")["total_matches"], 1)

    def test_block_quotation_excluded(self):
        self.assertEqual(row("> not only an example\n\nordinary prose", "contrast-not-only")["total_matches"], 0)

    def test_inline_prose_quote_remains_for_review(self):
        self.assertEqual(row('The note reads "not only".', "contrast-not-only")["total_matches"], 1)

    def test_bibliography_exclusion_and_resume(self):
        text = "# Main\nnot only\n\n## References\nnot only\n\n### Notes\nnot only\n\n## Appendix\nnot only"
        self.assertEqual(row(text, "contrast-not-only")["total_matches"], 2)
        self.assertEqual(row(text, "contrast-not-only", include_references=True)["total_matches"], 4)

    def test_mask_preserves_positions(self):
        text = "A `code` fragment.\n> quotation\n~~~\nexample\n~~~\nTail"
        masked = mask_markdown(text)
        self.assertEqual(len(masked), len(text))
        self.assertEqual([i for i,c in enumerate(masked) if c == "\n"], [i for i,c in enumerate(text) if c == "\n"])

    def test_plain_text_has_no_markdown_mask(self):
        self.assertEqual(row("`not only`", "contrast-not-only", markdown=False)["total_matches"], 1)

    def test_semicolon_citation_cluster(self):
        text = "The claim has support (Author A 2001; Author B 2002)."
        self.assertEqual(row(text, "citation-cluster")["total_matches"], 1)

    def test_wrapped_citation_cluster(self):
        text = "Support (Author A 2001;\nAuthor B 2002)."
        self.assertEqual(row(text, "citation-cluster")["total_matches"], 1)

    def test_numeric_semicolon_cluster(self):
        self.assertEqual(row("Support [1; 2].", "citation-cluster")["total_matches"], 1)

    def test_semicolon_without_citations(self):
        self.assertEqual(row("A pair (x; y) defines the relation.", "citation-cluster")["total_matches"], 0)

    def test_one_citation_is_not_a_cluster(self):
        self.assertEqual(row("The estimate changed (Author A 2001).", "citation-cluster")["total_matches"], 0)

    def test_parenthetical_size(self):
        text = "The claim (Author A 2001, Author B 2002, Author C 2003) is too compressed."
        self.assertEqual(row(text, "citation-size")["total_matches"], 1)

    def test_five_source_paragraph(self):
        text = " ".join(f"Author {chr(65+i)} ({2001+i}) identifies a problem." for i in range(5))
        self.assertEqual(row(text, "citation-density")["total_matches"], 1)

    def test_four_source_paragraph(self):
        text = " ".join(f"Author {chr(65+i)} ({2001+i}) identifies a problem." for i in range(4))
        self.assertEqual(row(text, "citation-density")["total_matches"], 0)

    def test_duplicate_source_counted_once(self):
        text = "Author A (2001) identifies a problem. " * 7
        self.assertEqual(row(text, "citation-density")["total_matches"], 0)

    def test_coauthors_count_as_one_source(self):
        text = "(Lane and Reed 2001; North et al. 2002; Quinn 2003; West 2004)"
        self.assertEqual(row(text, "citation-density")["total_matches"], 0)

    def test_paragraph_break_resets_source_count(self):
        text = "Author A (2001). Author B (2002). Author C (2003).\n\nAuthor D (2004). Author E (2005)."
        self.assertEqual(row(text, "citation-density")["total_matches"], 0)

    def test_numeric_citation_density(self):
        self.assertEqual(row("The claim has support [1-5].", "citation-density")["total_matches"], 1)
        self.assertEqual(row("The claim has support [1, 2, 3].", "citation-size")["total_matches"], 1)

    def test_numeric_range_is_bounded(self):
        self.assertEqual(row("An invalid list [1-999999999].", "citation-density")["total_matches"], 0)

    def test_author_parade(self):
        text = "Author A (2001) argues for it. Author B (2002) shows a difference. Author C (2003) adds another claim."
        self.assertEqual(row(text, "author-parade")["total_matches"], 1)

    def test_non_author_subjects_do_not_form_parade(self):
        text = "The record shows a date. The rubric shows a category. The table shows a total."
        self.assertEqual(row(text, "author-parade")["total_matches"], 0)

    def test_original_positive_examples_have_no_policy_matches(self):
        for rule in load("rules")["rules"]:
            with self.subTest(rule=rule["id"]):
                self.assertEqual(audit_text(rule["example"]["after"])["policy_matches"], 0)

    def test_network_not_used(self):
        with patch("socket.socket", side_effect=AssertionError("Network use")):
            result = audit_text("The catalog records each item's location.")
        self.assertEqual(result["total_matches"], 0)

    def test_read_bom_utf8_and_preserve_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.txt"
            before = b"\xef\xbb\xbfnot only a record\n"
            path.write_bytes(before)
            text, unit = read_manuscript(path)
            self.assertEqual(text, "not only a record\n")
            self.assertEqual(unit, "line")
            audit_text(text)
            self.assertEqual(path.read_bytes(), before)

    def test_invalid_utf8(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.txt"
            path.write_bytes(b"\xff\xfe")
            with self.assertRaises(ValueError):
                read_manuscript(path)

    def test_missing_file(self):
        with tempfile.TemporaryDirectory() as directory, self.assertRaises(ValueError):
            read_manuscript(Path(directory) / "absent.txt")

    def test_unsupported_extension(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.html"
            path.write_text("example")
            with self.assertRaises(ValueError):
                read_manuscript(path)

    def test_file_size_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.txt"
            path.write_text("1234")
            with patch("journalprose.review.MAX_INPUT_BYTES", 3), self.assertRaises(ValueError):
                read_manuscript(path)

    def test_corrupt_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.docx"
            path.write_text("invalid")
            with self.assertRaises(ValueError):
                read_manuscript(path)


PATTERN_CASES = {
    "contrast-not-but": "This is not a list but a framework.",
    "contrast-not-merely": "The value is not merely descriptive.",
    "contrast-not-only": "The record is not only complete.",
    "contrast-rather": "Use the register rather than the notes.",
    "contrast-instead": "Use the register instead of the notes.",
    "contrast-opposed": "Consider accuracy as opposed to completeness.",
    "contrast-no-longer": "The record is no longer incomplete but usable.",
    "translated-contrast": "\u4e0d\u662f\u5217\u8868\uff0c\u800c\u662f\u6846\u67b6\u3002",
    "purpose-template": "This paper argues that the record is incomplete.",
    "process-announcement": "In this section, we turn to the next issue.",
    "promotional-language": "A groundbreaking method appears here.",
    "shallow-ranking": "Completeness is deeper than accuracy.",
    "vague-reference": "This shows the value of the approach.",
    "stacked-hedge": "The estimate may perhaps differ.",
    "agreement-candidate": "These results is stable.",
    "article-mass-noun": "A evidence supports the claim.",
    "repeated-word": "The the record is complete.",
}


def make_pattern_test(check, text):
    def test(self):
        self.assertGreater(row(text, check)["total_matches"], 0)
    return test


for _check, _text in PATTERN_CASES.items():
    setattr(ReviewTests, "test_pattern_" + _check.replace("-", "_"), make_pattern_test(_check, _text))


try:
    import docx
except ImportError:
    docx = None


@unittest.skipIf(docx is None, "Optional python-docx is not installed")
class DocxTests(unittest.TestCase):
    def test_body_table_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.docx"
            d = docx.Document()
            d.add_paragraph("Before")
            table = d.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Left"
            table.cell(0, 1).text = "Right"
            d.add_paragraph("After")
            d.save(path)
            text, unit = read_manuscript(path)
            self.assertEqual(unit, "extracted_line")
            self.assertEqual(text.split("\n\n"), ["Before", "Left", "Right", "After"])

    def test_many_table_cells_are_retained(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.docx"
            d = docx.Document()
            table = d.add_table(rows=8, cols=4)
            expected = []
            for i, row_ in enumerate(table.rows):
                for j, cell in enumerate(row_.cells):
                    word = f"Cell-{i}-{j}"
                    expected.append(word)
                    cell.text = word
            d.save(path)
            self.assertEqual(read_manuscript(path)[0].split("\n\n"), expected)

    def test_merged_cells_are_not_duplicated(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.docx"
            d = docx.Document()
            table = d.add_table(rows=1, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
            d.save(path)
            self.assertEqual(read_manuscript(path)[0].count("Merged"), 1)

    def test_expanded_archive_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.docx"
            d = docx.Document()
            d.add_paragraph("Content")
            d.save(path)
            with patch("journalprose.review.MAX_DOCX_EXPANDED_BYTES", 10), self.assertRaises(ValueError):
                read_manuscript(path)


if __name__ == "__main__":
    unittest.main()
